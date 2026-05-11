from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree

# ── Colour palette ──────────────────────────────────────────────────────────
DARK_BG   = RGBColor(0x0F, 0x17, 0x2A)   # #0f172a  navy sidebar
MID_BG    = RGBColor(0x1E, 0x29, 0x3B)   # #1e293b
ACCENT    = RGBColor(0x25, 0x63, 0xEB)   # #2563eb  blue
ACCENT2   = RGBColor(0x60, 0xA5, 0xFA)   # #60a5fa  light blue
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF1, 0xF5, 0xF9)
MUTED     = RGBColor(0x94, 0xA3, 0xB8)
CRITICAL  = RGBColor(0xDC, 0x26, 0x26)
HIGH      = RGBColor(0xEA, 0x58, 0x0C)
MEDIUM    = RGBColor(0xD9, 0x77, 0x06)
LOW       = RGBColor(0x16, 0xA3, 0x4A)
JIRA_BLUE = RGBColor(0x1D, 0x4E, 0xD8)
CONF_PURP = RGBColor(0x6D, 0x28, 0xD9)
SNOW_GRN  = RGBColor(0x15, 0x80, 0x3D)

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

BLANK = prs.slide_layouts[6]  # completely blank layout


# ══════════════════════════════════════════════════════════════════════════════
# Helper utilities
# ══════════════════════════════════════════════════════════════════════════════

def bg(slide, color=DARK_BG):
    """Fill slide background with a solid colour."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(slide, l, t, w, h, fill=None, line=None, line_w=Pt(0)):
    """Add a filled rectangle."""
    sh = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.line.fill.background()
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = line
        sh.line.width = line_w
    else:
        sh.line.fill.background()
    return sh

def txt(slide, text, l, t, w, h, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, wrap=True):
    """Add a text box."""
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txb

def pill(slide, label, l, t, w, h, bg_c, fg_c, size=10, bold=True):
    """Rounded pill / badge."""
    sh = slide.shapes.add_shape(5, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = bg_c
    sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = False
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.CENTER
    run = para.add_run(); run.text = label
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = fg_c; run.font.name = "Calibri"
    return sh

def divider(slide, t, color=MID_BG):
    box(slide, 0, t, 13.33, 0.02, fill=color)

def section_label(slide, text, l=0.45, t=0.18):
    txt(slide, text.upper(), l, t, 4, 0.3, size=9, bold=True,
        color=MUTED, align=PP_ALIGN.LEFT)

def slide_title(slide, title, subtitle=None, tl=0.45, tt=0.52):
    txt(slide, title, tl, tt, 12, 0.6, size=28, bold=True, color=WHITE)
    if subtitle:
        txt(slide, subtitle, tl, tt+0.58, 12, 0.4, size=14, color=MUTED)

def kpi_card(slide, l, t, w, h, value, label, val_color, accent_color):
    box(slide, l, t, w, h, fill=MID_BG, line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
    box(slide, l, t, w, 0.055, fill=accent_color)   # top stripe
    txt(slide, value, l+0.12, t+0.15, w-0.24, 0.55, size=32, bold=True, color=val_color)
    txt(slide, label, l+0.12, t+0.72, w-0.24, 0.35, size=10, color=MUTED)

def flow_arrow(slide, x, y, horiz=True):
    """Simple right-pointing or down-pointing arrow label."""
    ch = "→" if horiz else "↓"
    txt(slide, ch, x, y, 0.35, 0.35, size=18, color=MUTED, align=PP_ALIGN.CENTER)

def icon_card(slide, l, t, w, h, icon, title, body, accent):
    box(slide, l, t, w, h, fill=MID_BG, line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
    box(slide, l, t, 0.06, h, fill=accent)  # left stripe
    txt(slide, icon,  l+0.15, t+0.12, 0.5, 0.45, size=20, color=accent)
    txt(slide, title, l+0.65, t+0.10, w-0.80, 0.35, size=13, bold=True, color=WHITE)
    txt(slide, body,  l+0.65, t+0.42, w-0.80, h-0.55, size=10, color=MUTED, wrap=True)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
# gradient accent bar left
box(s, 0, 0, 0.5, 7.5, fill=ACCENT)
box(s, 0.5, 0, 0.08, 7.5, fill=RGBColor(0x1D,0x4E,0xD8))

txt(s, "SecureOps", 1.1, 1.6, 10, 1.1, size=52, bold=True, color=WHITE)
txt(s, "Vulnerability & Incident Management Platform",
    1.1, 2.75, 11, 0.7, size=22, color=ACCENT2)
divider(s, 3.55, ACCENT)
txt(s, "Automated security scanning  ·  Incident lifecycle management  ·  Jira / ServiceNow / Confluence integration",
    1.1, 3.7, 11.5, 0.5, size=13, color=MUTED)
txt(s, "Platform Engineering  ·  May 2026",
    1.1, 6.7, 6, 0.4, size=11, color=MUTED)
pill(s, "CONFIDENTIAL", 10.8, 6.7, 1.8, 0.35,
     RGBColor(0x2D,0x3F,0x55), MUTED, size=9)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Problem Statement
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Context")
slide_title(s, "The Problem We're Solving")

problems = [
    ("🔍", "Scan results buried in CI logs",
     "OWASP, SpotBugs, Trivy, and SonarQube results exist in separate pipeline artefacts with no unified view."),
    ("⏱", "Manual triage wastes engineer time",
     "Engineers must SSH into CI, download JSON reports, and manually decide which findings need action."),
    ("🔗", "No link between findings and tickets",
     "Security findings are not automatically tracked in Jira or ServiceNow — critical CVEs can go unassigned for days."),
    ("📄", "RCA documentation is ad-hoc",
     "Post-incident write-ups live in personal docs or email threads, never in a searchable Confluence space."),
    ("🚨", "FCA/PRA compliance risk",
     "Regulated firms must demonstrate timely remediation of critical vulnerabilities. Manual processes create audit gaps."),
]

for i, (icon, title, body) in enumerate(problems):
    col = i % 2
    row = i // 2
    lft = 0.45 + col * 6.55
    top = 1.55 + row * 1.55
    w = 6.1
    icon_card(s, lft, top, w, 1.4, icon, title, body, CRITICAL if i==4 else ACCENT)

# bottom bar
box(s, 0, 6.9, 13.33, 0.6, fill=MID_BG)
txt(s, "Result: critical CVEs missed, SLA breaches, manual overhead, and compliance exposure.",
    0.45, 6.97, 12, 0.4, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Solution Overview
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Solution")
slide_title(s, "SecureOps — What It Does")

pillars = [
    (ACCENT,   "01", "Unified\nVulnerability View",
     "Aggregates findings from OWASP, SpotBugs, Trivy & SonarQube into one filterable dashboard"),
    (CRITICAL, "02", "Automated\nIncident Creation",
     "Critical/High findings automatically create incidents with SLA countdown timers and triage queue"),
    (JIRA_BLUE,"03", "Jira & ServiceNow\nIntegration",
     "Each incident auto-opens a Jira SEC ticket and ServiceNow INC record — no manual entry"),
    (CONF_PURP,"04", "Confluence RCA\nGeneration",
     "Structured RCA pages (timeline, root cause, action items, prevention) created automatically per incident"),
    (LOW,      "05", "Cross-linked\nAudit Trail",
     "Every finding links to its incident, Jira ticket, ServiceNow record, and Confluence RCA — full traceability"),
]

for i, (color, num, title, body) in enumerate(pillars):
    l = 0.45 + i * 2.52
    box(s, l, 1.55, 2.3, 4.5, fill=MID_BG, line=color, line_w=Pt(2))
    box(s, l, 1.55, 2.3, 0.07, fill=color)
    txt(s, num, l+0.15, 1.65, 0.6, 0.5, size=22, bold=True, color=color)
    txt(s, title, l+0.15, 2.25, 2.0, 0.7, size=13, bold=True, color=WHITE)
    txt(s, body,  l+0.15, 2.98, 2.0, 2.8, size=10, color=MUTED, wrap=True)

box(s, 0, 6.6, 13.33, 0.9, fill=MID_BG)
txt(s, "One platform — from pipeline scan result to resolved, documented incident.",
    0.45, 6.72, 12, 0.4, size=13, bold=True, color=ACCENT2, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Architecture
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Architecture")
slide_title(s, "How It All Connects")

# Left column — Pipeline tools
box(s, 0.35, 1.4, 2.7, 5.5, fill=RGBColor(0x0A,0x10,0x1E),
    line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
txt(s, "Pipeline Tools", 0.35, 1.4, 2.7, 0.4, size=10, bold=True,
    color=MUTED, align=PP_ALIGN.CENTER)
tools = [
    (HIGH,    "OWASP Dep-Check"),
    (ACCENT,  "SpotBugs / SAST"),
    (MEDIUM,  "Trivy Container"),
    (LOW,     "SonarQube"),
    (MUTED,   "GitLab CI Webhook"),
]
for i, (c, name) in enumerate(tools):
    box(s, 0.5, 1.95 + i*0.85, 2.4, 0.65, fill=MID_BG,
        line=c, line_w=Pt(1))
    txt(s, name, 0.7, 2.05 + i*0.85, 2.1, 0.4, size=11, color=WHITE)

# Arrow
flow_arrow(s, 3.1, 3.8)

# Middle — Backend API
box(s, 3.5, 1.4, 3.1, 5.5, fill=RGBColor(0x0A,0x10,0x1E),
    line=ACCENT, line_w=Pt(2))
txt(s, "SecureOps Backend API", 3.5, 1.4, 3.1, 0.4, size=10, bold=True,
    color=ACCENT2, align=PP_ALIGN.CENTER)
backend = [
    (ACCENT,  "Webhook receiver"),
    (ACCENT,  "Report parser"),
    (ACCENT,  "Findings store (PostgreSQL)"),
    (ACCENT,  "Incident engine"),
    (ACCENT,  "REST API for dashboard"),
    (ACCENT,  "Integration dispatcher"),
]
for i, (c, name) in enumerate(backend):
    box(s, 3.65, 1.95 + i*0.80, 2.8, 0.60, fill=MID_BG, line=c, line_w=Pt(1))
    txt(s, name, 3.82, 2.05 + i*0.80, 2.55, 0.38, size=10, color=WHITE)

# Arrow
flow_arrow(s, 6.7, 3.8)

# Right — Consumers
box(s, 7.1, 1.4, 2.6, 2.5, fill=RGBColor(0x0A,0x10,0x1E),
    line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
txt(s, "SecureOps Dashboard", 7.1, 1.4, 2.6, 0.4, size=10, bold=True,
    color=MUTED, align=PP_ALIGN.CENTER)
dash_items = ["Vulnerability table", "Triage queue", "Incident view", "Charts / KPIs"]
for i, name in enumerate(dash_items):
    txt(s, f"• {name}", 7.25, 1.9 + i*0.48, 2.3, 0.38, size=10, color=WHITE)

box(s, 7.1, 4.1, 2.6, 2.8, fill=RGBColor(0x0A,0x10,0x1E),
    line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
txt(s, "External Systems", 7.1, 4.1, 2.6, 0.4, size=10, bold=True,
    color=MUTED, align=PP_ALIGN.CENTER)
ext_items = [
    (JIRA_BLUE, "Jira (SEC project)"),
    (SNOW_GRN,  "ServiceNow (INC)"),
    (CONF_PURP, "Confluence (RCA)"),
    (MUTED,     "Slack alerts"),
]
for i, (c, name) in enumerate(ext_items):
    box(s, 7.25, 4.6 + i*0.54, 2.3, 0.42, fill=MID_BG, line=c, line_w=Pt(1))
    txt(s, name, 7.42, 4.68 + i*0.54, 2.1, 0.3, size=10, color=WHITE)

# Far right — Data flow note
box(s, 10.0, 1.4, 3.0, 5.5, fill=RGBColor(0x0A,0x10,0x1E),
    line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
txt(s, "Data Flow", 10.0, 1.4, 3.0, 0.4, size=10, bold=True,
    color=MUTED, align=PP_ALIGN.CENTER)
steps = [
    "① Push code to GitLab",
    "② Pipeline runs all scans",
    "③ CI webhook fires to API",
    "④ API downloads artefacts",
    "⑤ Reports parsed & stored",
    "⑥ Auto-incidents for Crit/High",
    "⑦ Jira + SNOW + Confluence",
    "⑧ Dashboard refreshes live",
]
for i, step in enumerate(steps):
    txt(s, step, 10.1, 1.92 + i*0.60, 2.8, 0.48, size=10,
        color=WHITE if i % 2 == 0 else MUTED)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — Security Scanning Layer
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Security Scanning")
slide_title(s, "Four-Layer Scan Coverage in Every Pipeline Run")

scans = [
    (HIGH,    "OWASP\nDependency-Check",
     "What it catches",
     "Known CVEs in Maven dependencies\nCVSS score, affected version, fix version",
     "Gate threshold",
     "Build fails on CVSS ≥ 7.0\nWeekly audit uses CVSS ≥ 4.0"),
    (ACCENT,  "SpotBugs +\nfind-sec-bugs",
     "What it catches",
     "SQL injection, hardcoded secrets,\nunsafe deserialization, CSRF gaps",
     "Gate threshold",
     "Medium+ confidence issues\nfail the build"),
    (MEDIUM,  "Trivy\nContainer Scan",
     "What it catches",
     "OS package CVEs in base image,\nJAR CVEs in container layers",
     "Gate threshold",
     "Critical/High image CVEs\nblock the deploy stage"),
    (LOW,     "SonarQube\nQuality Gate",
     "What it catches",
     "Code smells, complexity, duplicates,\nsecurity hotspots, coverage drops",
     "Gate threshold",
     "Rating A/A/A required\n80% line coverage minimum"),
]

for i, (color, name, l1, v1, l2, v2) in enumerate(scans):
    lft = 0.45 + i * 3.22
    box(s, lft, 1.5, 3.0, 5.3, fill=MID_BG,
        line=color, line_w=Pt(2))
    box(s, lft, 1.5, 3.0, 0.07, fill=color)
    txt(s, name,  lft+0.15, 1.6, 2.7, 0.65, size=14, bold=True, color=WHITE)
    txt(s, l1, lft+0.15, 2.42, 2.7, 0.28, size=9, bold=True, color=color)
    txt(s, v1, lft+0.15, 2.72, 2.7, 1.0,  size=10, color=OFF_WHITE, wrap=True)
    txt(s, l2, lft+0.15, 3.85, 2.7, 0.28, size=9, bold=True, color=color)
    txt(s, v2, lft+0.15, 4.15, 2.7, 1.0,  size=10, color=OFF_WHITE, wrap=True)
    # outcome badge
    pill(s, "AUTO-INCIDENT ON FAILURE", lft+0.25, 6.35, 2.5, 0.32,
         RGBColor(0x1E,0x29,0x3B), color, size=8)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — Incident Lifecycle
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Incident Management")
slide_title(s, "End-to-End Incident Lifecycle")

stages = [
    (CRITICAL,  "1\nDetect",   "Pipeline scan\nfinds Critical/High\nvulnerability"),
    (HIGH,      "2\nCreate",   "SecureOps auto-\ncreates incident\nwith SLA timer"),
    (JIRA_BLUE, "3\nJira",     "SEC-xxx ticket\nopened with full\nfinding detail"),
    (SNOW_GRN,  "4\nServiceNow","INC-xxxxxxx raised,\nassigned to\nPlatform Security"),
    (CONF_PURP, "5\nConfluence","RCA page auto-\ngenerated in\nPLATFORM space"),
    (MEDIUM,    "6\nResolve",  "Fix merged,\nincident closed,\nRCA published"),
]

for i, (color, stage, desc) in enumerate(stages):
    lft = 0.5 + i * 2.06
    # circle
    circ = slide.shapes if False else None
    box(s, lft+0.6, 1.55, 0.9, 0.9, fill=color)  # square stand-in for circle
    txt(s, stage, lft+0.55, 1.55, 1.0, 0.9, size=11, bold=True,
        color=WHITE, align=PP_ALIGN.CENTER)
    # connector arrow (not for last)
    if i < 5:
        txt(s, "→", lft+1.6, 1.82, 0.4, 0.4, size=20, color=MUTED,
            align=PP_ALIGN.CENTER)
    # description card
    box(s, lft+0.1, 2.7, 1.8, 1.4, fill=MID_BG,
        line=color, line_w=Pt(1))
    txt(s, desc, lft+0.18, 2.82, 1.65, 1.2, size=10, color=OFF_WHITE, wrap=True)

# SLA table
box(s, 0.45, 4.4, 12.4, 2.7, fill=RGBColor(0x0A,0x10,0x1E),
    line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
txt(s, "SLA Policy", 0.45, 4.4, 12.4, 0.38, size=10, bold=True,
    color=MUTED, align=PP_ALIGN.CENTER)

sla_data = [
    ("Severity", "Incident Creation SLA", "Remediation SLA",  "Auto-escalate"),
    ("Critical",  "Within 2 hours",         "48 hours",          "At 1 hour breach"),
    ("High",      "Within 4 hours",         "7 days",            "At 3 hour breach"),
    ("Medium",    "Within 24 hours",        "30 days",           "At 12 hour breach"),
    ("Low",       "Within 72 hours",        "90 days",           "No auto-escalation"),
]
col_w = [2.0, 3.2, 3.2, 3.2]
col_x = [0.55, 2.65, 5.95, 9.25]
row_colors = [MUTED, CRITICAL, HIGH, MEDIUM, LOW]
for r, row in enumerate(sla_data):
    for c, cell in enumerate(row):
        fc = WHITE if r == 0 else OFF_WHITE
        bold = r == 0
        txt(s, cell, col_x[c], 4.85 + r*0.42, col_w[c], 0.38,
            size=10, color=fc if r > 0 else row_colors[r], bold=bold)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — Integration Detail
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Integrations")
slide_title(s, "Jira · ServiceNow · Confluence — Automatic on Every Incident")

# Three integration cards
integ = [
    (JIRA_BLUE, "Jira", "SEC Project",
     [("Ticket type", "Bug / Security"),
      ("Priority",    "Maps from CVE severity"),
      ("Labels",      "security, cve, pipeline"),
      ("Description", "Finding detail + CVSS + affected component"),
      ("Links",       "ServiceNow INC + Confluence RCA"),
      ("Status sync", "Closed when incident resolved")]),
    (SNOW_GRN, "ServiceNow", "Platform Security Queue",
     [("Category",    "Information Security"),
      ("Urgency",     "1-Critical, 2-High, 3-Medium"),
      ("Impact",      "Derived from CVSS scope"),
      ("Assignment",  "Platform Security group"),
      ("Description", "Full finding + Jira ticket link"),
      ("Resolution",  "Auto-updated on close")]),
    (CONF_PURP, "Confluence", "PLATFORM Space",
     [("Page title",  "RCA-YYYY-MM-DD-[incident]"),
      ("Sections",    "Summary, Timeline, Root Cause"),
      ("Action items","Linked Jira tickets per item"),
      ("Prevention",  "Automated recommendations"),
      ("Cross-links", "Jira SEC-xxx + ServiceNow INC"),
      ("Status",      "Draft → In Review → Published")]),
]

for i, (color, title, sub, rows) in enumerate(integ):
    lft = 0.45 + i * 4.3
    box(s, lft, 1.48, 4.0, 5.5, fill=MID_BG,
        line=color, line_w=Pt(2))
    box(s, lft, 1.48, 4.0, 0.07, fill=color)
    # Header
    pill(s, title, lft+0.15, 1.58, 1.2, 0.38, color, WHITE, size=12, bold=True)
    txt(s, sub, lft+1.45, 1.62, 2.4, 0.32, size=10, color=MUTED)
    # Rows
    for j, (key, val) in enumerate(rows):
        row_top = 2.15 + j * 0.75
        box(s, lft+0.15, row_top, 3.7, 0.62, fill=RGBColor(0x0F,0x17,0x2A),
            line=RGBColor(0x2D,0x3F,0x55), line_w=Pt(1))
        txt(s, key, lft+0.28, row_top+0.06, 1.1, 0.25, size=8, bold=True, color=color)
        txt(s, val, lft+0.28, row_top+0.30, 3.45, 0.25, size=9, color=OFF_WHITE)

# Slack note at bottom
box(s, 0.45, 7.05, 12.4, 0.38, fill=RGBColor(0x0A,0x10,0x1E))
txt(s, "All integrations also trigger a colour-coded Slack alert to #platform-alerts with direct links to Jira and Confluence.",
    0.55, 7.1, 12.2, 0.28, size=9, color=MUTED, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — Dashboard Walkthrough
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Dashboard")
slide_title(s, "SecureOps UI — Six Core Views")

pages = [
    (ACCENT,   "Dashboard",        "KPI cards · trend chart · severity donut · scan coverage · live activity feed"),
    (CRITICAL, "All Findings",     "Filterable table of every scan finding — OWASP, SAST, Container, SonarQube"),
    (HIGH,     "Triage Queue",     "Unactioned findings with SLA countdown timers and one-click Create Incident"),
    (MEDIUM,   "Incidents",        "Active + resolved incidents with Jira, ServiceNow, and Confluence columns"),
    (CONF_PURP,"RCA & Reports",    "Confluence RCA page list + inline preview + weekly bar chart"),
    (JIRA_BLUE,"Jira / Confluence","All auto-created tickets and pages with status, assignee, and cross-links"),
]

for i, (color, page, desc) in enumerate(pages):
    row = i // 2
    col = i %  2
    lft = 0.45 + col * 6.45
    top = 1.5  + row * 1.7
    box(s, lft, top, 6.1, 1.55, fill=MID_BG,
        line=color, line_w=Pt(2))
    box(s, lft, top, 0.07, 1.55, fill=color)
    txt(s, f"0{i+1}", lft+0.25, top+0.12, 0.55, 0.45, size=20, bold=True, color=color)
    txt(s, page,      lft+0.82, top+0.12, 4.8,  0.4,  size=14, bold=True, color=WHITE)
    txt(s, desc,      lft+0.82, top+0.60, 4.9,  0.75, size=10, color=MUTED, wrap=True)

box(s, 0, 6.75, 13.33, 0.75, fill=MID_BG)
txt(s, "Each incident row opens a slide-in detail panel: finding details, full timeline, Jira/SNOW/Confluence links, and RCA summary.",
    0.45, 6.87, 12.4, 0.38, size=11, color=OFF_WHITE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — Tech Stack
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Technology")
slide_title(s, "Tech Stack")

layers = [
    ("Application", ACCENT, [
        ("Java 17 / Spring Boot 3.2", "Core application framework"),
        ("Maven",                     "Build, plugin orchestration, dependency management"),
        ("JaCoCo",                    "Code coverage — 80% line / 75% branch gates"),
    ]),
    ("Security Scanning", CRITICAL, [
        ("OWASP Dependency-Check",    "CVE scanning of Maven dependencies (CVSS gate ≥ 7)"),
        ("SpotBugs + find-sec-bugs",  "SAST — SQL injection, secrets, deserialization"),
        ("Trivy",                     "Container image vulnerability scanning"),
        ("SonarQube Enterprise",      "Unified quality gate — ratings A/A/A, 80% coverage"),
    ]),
    ("Pipeline & Deployment", MEDIUM, [
        ("GitLab CI",                 "9-stage pipeline: validate → build → test → security → deploy"),
        ("ArgoCD + Helm",             "GitOps deployment to EKS across dev / staging / prod"),
        ("AWS ECR",                   "Immutable-tag container registry with KMS encryption"),
        ("Terraform",                 "IaC for EKS, ECR, IAM — isolated state per environment"),
    ]),
    ("Incident & Notifications", CONF_PURP, [
        ("Python notify.py",          "Unified dispatcher — Jira, ServiceNow, Confluence, Slack"),
        ("Jira REST API v3",          "Auto-creates SEC-xxx Bug tickets with structured ADF body"),
        ("ServiceNow Table API",      "Raises and closes INC records with urgency/impact mapping"),
        ("Confluence REST API v2",    "Creates structured RCA pages in PLATFORM space"),
    ]),
]

for li, (layer_name, color, items) in enumerate(layers):
    col = li % 2
    row = li // 2
    lft = 0.45 + col * 6.45
    top = 1.45 + row * 2.85
    box(s, lft, top, 6.1, 2.65, fill=MID_BG,
        line=color, line_w=Pt(1))
    box(s, lft, top, 6.1, 0.07, fill=color)
    txt(s, layer_name, lft+0.15, top+0.12, 5.7, 0.38, size=12, bold=True, color=color)
    for j, (tech, desc) in enumerate(items):
        box(s, lft+0.15, top+0.60 + j*0.48, 5.8, 0.40,
            fill=RGBColor(0x0F,0x17,0x2A))
        txt(s, tech, lft+0.28, top+0.65 + j*0.48, 2.2, 0.28, size=10, bold=True, color=WHITE)
        txt(s, desc, lft+2.55, top+0.65 + j*0.48, 3.45, 0.28, size=9, color=MUTED)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — Benefits / ROI
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Business Case")
slide_title(s, "Why This Matters")

metrics = [
    (CRITICAL,  "< 2h",     "Critical incident\ncreation SLA"),
    (HIGH,      "18h",      "Mean time\nto resolve"),
    (ACCENT,    "0",        "Manual ticket\ncreation steps"),
    (LOW,       "100%",     "Audit trail\ncoverage"),
]

for i, (color, val, label) in enumerate(metrics):
    lft = 0.45 + i * 3.22
    kpi_card(s, lft, 1.45, 3.0, 1.65, val, label, color, color)

benefits = [
    (LOW,      "FCA / PRA Compliance",
     "Timestamped, cross-linked incident records satisfy regulatory requirements for vulnerability tracking and remediation evidence."),
    (ACCENT,   "Zero Manual Overhead",
     "From CVE detection to Jira ticket to Confluence RCA — the entire workflow executes automatically within seconds of pipeline completion."),
    (CONF_PURP,"Institutional Knowledge",
     "Every incident produces a permanent, searchable Confluence RCA — eliminating tribal knowledge and enabling faster future triage."),
    (HIGH,     "Earlier Detection",
     "CVSS ≥ 4 weekly audit catches medium-severity issues before they reach critical status in production."),
    (JIRA_BLUE,"Single Pane of Glass",
     "Engineers no longer trawl CI logs. All findings, incidents, tickets, and RCAs are visible in one portal."),
    (MEDIUM,   "Scalable to All Java Services",
     "The same pipeline template and notification dispatcher can be applied to every Java microservice in the estate."),
]

for i, (color, title, body) in enumerate(benefits):
    col = i % 2
    row = i // 2
    lft = 0.45 + col * 6.45
    top = 3.35 + row * 1.3
    box(s, lft, top, 6.1, 1.18, fill=MID_BG,
        line=color, line_w=Pt(1))
    box(s, lft, top, 0.06, 1.18, fill=color)
    txt(s, title, lft+0.22, top+0.08, 5.6, 0.32, size=12, bold=True, color=WHITE)
    txt(s, body,  lft+0.22, top+0.45, 5.7, 0.65, size=10, color=MUTED, wrap=True)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — Roadmap / Next Steps
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
section_label(s, "Roadmap")
slide_title(s, "Next Steps")

phases = [
    ("Phase 1\nNow — Week 2",   ACCENT,   [
        "Deploy backend API to dev EKS namespace",
        "Wire GitLab CI webhook → backend",
        "Connect SonarQube API poller",
        "Validate end-to-end: scan → incident → Jira",
    ]),
    ("Phase 2\nWeeks 3–4",      HIGH,     [
        "ServiceNow live integration (non-prod)",
        "Confluence RCA auto-publish to PLATFORM",
        "Slack webhook for #platform-alerts",
        "Dashboard served from internal URL",
    ]),
    ("Phase 3\nWeeks 5–6",      LOW,      [
        "Prod cutover — all Java pipelines onboarded",
        "SLA reporting to management (weekly email)",
        "DORA metrics dashboard tab",
        "Compliance export for FCA audit pack",
    ]),
    ("Phase 4\nFuture",         CONF_PURP,[
        "Extend to Node.js / Python services",
        "ML-assisted false-positive suppression",
        "Automated remediation PRs via Renovate",
        "Integration with Veracode / Snyk",
    ]),
]

for i, (phase, color, items) in enumerate(phases):
    lft = 0.45 + i * 3.22
    box(s, lft, 1.45, 3.0, 5.5, fill=MID_BG, line=color, line_w=Pt(2))
    box(s, lft, 1.45, 3.0, 0.07, fill=color)
    txt(s, phase, lft+0.15, 1.55, 2.7, 0.7, size=12, bold=True, color=color)
    for j, item in enumerate(items):
        box(s, lft+0.15, 2.4 + j*1.1, 2.7, 0.9,
            fill=RGBColor(0x0F,0x17,0x2A))
        txt(s, f"✓ {item}", lft+0.28, 2.5 + j*1.1, 2.5, 0.75,
            size=10, color=OFF_WHITE, wrap=True)

box(s, 0, 7.1, 13.33, 0.4, fill=ACCENT)
txt(s, "Phase 1 can begin immediately — no new infrastructure required beyond the backend service deployment.",
    0.45, 7.16, 12.4, 0.26, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 12 — Closing
# ══════════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
bg(s)
box(s, 0, 0, 0.5, 7.5, fill=ACCENT)
box(s, 0.5, 0, 0.08, 7.5, fill=RGBColor(0x1D,0x4E,0xD8))

txt(s, "Thank You", 1.1, 1.6, 10, 1.0, size=52, bold=True, color=WHITE)
divider(s, 2.8, ACCENT)

summary_points = [
    "✓  Unified view of all pipeline security findings",
    "✓  Automatic incident creation with full SLA management",
    "✓  Jira + ServiceNow + Confluence integrated — zero manual steps",
    "✓  Permanent, auditable RCA trail for every incident",
    "✓  Built on the existing Java / GitLab CI / AWS stack",
]
for i, point in enumerate(summary_points):
    txt(s, point, 1.1, 3.1 + i*0.62, 10, 0.48, size=13, color=OFF_WHITE)

txt(s, "Questions & Discussion", 1.1, 6.3, 8, 0.45, size=16, bold=True, color=ACCENT2)
txt(s, "Platform Engineering  ·  May 2026  ·  CONFIDENTIAL",
    1.1, 6.88, 10, 0.35, size=10, color=MUTED)


# ══════════════════════════════════════════════════════════════════════════════
# Save
# ══════════════════════════════════════════════════════════════════════════════
out = r"C:\Users\dilip\java-secure-pipeline\SecureOps-Platform-Overview.pptx"
prs.save(out)
print(f"Saved: {out}")
